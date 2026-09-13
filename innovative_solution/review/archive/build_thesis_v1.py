from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PACKAGE = ROOT / "innovative_solution"
OUTPUT = PACKAGE / "outputs"
PLOTS = OUTPUT / "plots"
THESIS_DIR = PACKAGE / "thesis"
DOCX_PATH = THESIS_DIR / "基于机器学习的TFT-LCD多工序质量特性预测研究_论文初稿.docx"


def load_json(name: str) -> dict:
    return json.loads((OUTPUT / name).read_text(encoding="utf-8"))


def set_run_font(run, chinese: str = "宋体", western: str = "Times New Roman", size=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9)


def add_native_equation(doc: Document, text: str, number: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justify = OxmlElement("m:jc")
    justify.set(qn("m:val"), "center")
    math_para_pr.append(justify)
    math_para.append(math_para_pr)
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    paragraph._p.append(math_para)
    number_run = paragraph.add_run(f"    ({number})")
    set_run_font(number_run, size=11)


def add_body(doc: Document, text: str, *, indent=True, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if indent else None
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_heading(doc: Document, text: str, level: int):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_figure(doc: Document, filename: str, caption: str, width=6.15):
    path = PLOTS / filename
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths=None):
    caption_paragraph = add_caption(doc, caption)
    caption_paragraph.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(header))
        set_run_font(run, chinese="黑体", size=10, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                shade_cell(cell, "F4F8FB")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.5)
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.3)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_specs = {
        1: (16, "黑体", True, True),
        2: (14, "黑体", True, False),
        3: (12, "黑体", True, False),
    }
    for level, (size, font, bold, page_break) in heading_specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.paragraph_format.space_before = Pt(10 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = page_break

    add_page_number(section.footer.paragraphs[0])


def chapter_model_table(summary: pd.DataFrame) -> list[list[str]]:
    scopes = ["random_oof", "rolling_oof", "temporal_holdout"]
    models = ["Dummy", "Ridge", "ExtraTrees", "XGBoost", "XGB-Recency", "MLP-Ensemble", "RobustBlend"]
    rows = []
    for model in models:
        values = []
        for scope in scopes:
            match = summary.loc[(summary.model == model) & (summary.scope == scope), "mse"]
            values.append(f"{float(match.iloc[0]):.6f}" if len(match) else "-")
        holdout_r2 = summary.loc[
            (summary.model == model) & (summary.scope == "temporal_holdout"), "r2"
        ]
        rows.append([model, *values, f"{float(holdout_r2.iloc[0]):.4f}" if len(holdout_r2) else "-"])
    return rows


def build() -> Path:
    THESIS_DIR.mkdir(parents=True, exist_ok=True)
    base = load_json("run_manifest.json")
    advanced = load_json("advanced_summary.json")
    phase_a = load_json("phase_a_manifest.json")
    phase_b = load_json("phase_b_manifest.json")
    model_summary = pd.read_csv(OUTPUT / "model_summary.csv")
    feature_audit = pd.read_csv(OUTPUT / "feature_audit.csv")
    operation_importance = pd.read_csv(OUTPUT / "operation_importance.csv")
    operation_drift = pd.read_csv(OUTPUT / "operation_drift.csv")
    reducer_summary = pd.read_csv(OUTPUT / "transform_reducer_summary.csv")
    ablation = pd.read_csv(OUTPUT / "ablation_summary.csv")
    residual_tests = pd.read_csv(OUTPUT / "residual_statistical_tests.csv")
    distribution = pd.read_csv(OUTPUT / "theoretical_distribution_fit.csv")
    row_missing = pd.read_csv(OUTPUT / "row_missing_report.csv")

    audit_counts = feature_audit.status.value_counts().to_dict()
    tested_features = int(distribution.loc[distribution.best_by_aic.astype(bool), "feature"].nunique())
    consistent_features = int(
        distribution.loc[distribution.best_by_aic.astype(bool), "distribution_consistent_at_5pct"]
        .astype(bool)
        .sum()
    )
    top_operations = operation_importance.head(5).operation.astype(str).tolist()
    drift_operations = operation_drift.head(5).operation.astype(str).tolist()
    y_mean = float(row_missing.Y.mean())
    y_std = float(row_missing.Y.std(ddof=1))
    y_min = float(row_missing.Y.min())
    y_max = float(row_missing.Y.max())

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "基于机器学习的TFT-LCD多工序质量特性预测研究"
    doc.core_properties.subject = "智能制造质量预测论文初稿"
    doc.core_properties.keywords = "TFT-LCD 机器学习 多工序 质量特性预测 数据漂移"
    doc.core_properties.author = ""

    cover = doc.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Cm(5.0)
    cover.add_run("基于机器学习的TFT-LCD多工序质量特性预测研究")
    cover_ppr = cover._p.get_or_add_pPr()
    cover_border = cover_ppr.find(qn("w:pBdr"))
    if cover_border is not None:
        cover_ppr.remove(cover_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Cm(2.2)
    run = subtitle.add_run("论文初稿  第一版")
    set_run_font(run, chinese="黑体", size=16)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Cm(6.0)
    run = date.add_run("2026年9月")
    set_run_font(run, size=13)

    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("摘  要")
    set_run_font(run, chinese="黑体", size=16, bold=True)
    add_body(
        doc,
        f"TFT-LCD制造包含数百道工序，生产参数具有维度高、样本少、量纲混合、缺失模式复杂和批次分布变化明显等特点。传统抽样检测难以及时覆盖全部产品，因此本文以阿里云天池智能制造质量预测复赛数据为对象，研究利用设备、工艺参数与时间信息预测连续质量特性值Y的方法。原始训练集包含{base['train_rows']}条样本和5954个字段，测试A和测试B分别包含{base['test_a_rows']}条和{base['test_b_rows']}条样本。",
    )
    add_body(
        doc,
        f"本文建立工序感知且面向批次漂移的机器学习流程。数据处理阶段逐列记录删除和转换依据，采用工序对应设备组均值、训练折中位数和缺失指示完成缺失建模；对时间字段构造工序起止、跨度及相邻工序间隔；对不同量纲变量先进行中位数和四分位距标准化，再生成工序级稳健统计。特征选择联合ExtraTrees贡献、目标相关性和前后批次漂移，从{base['clean_numeric_features']}个清洗后数值变量中稳定选择关键变量。模型阶段比较Ridge、RandomForest、ExtraTrees、XGBoost、近期加权XGBoost和轻量MLP，并采用随机交叉验证、滚动时间验证和时间外推评估区分同分布拟合能力与跨批次泛化能力。",
    )
    add_body(
        doc,
        f"实验表明，随机OOF上的XGBoost均方误差为0.010063，而时间外推误差为0.032888，说明随机切分明显低估了未来批次风险。基础非负融合在时间外推集上的MSE为{base['holdout_blend_metrics']['mse']:.6f}。A阶段仅使用原训练标签，通过三个扩展窗口和三随机种子模型选择均匀权重XGBoost；在策略锁定后进行公开答案评价，A集MSE由0.030016降至{phase_a['selected_external_a_mse']:.6f}，R²为{phase_a['selected_external_a_r2']:.4f}。A答案公布后，B阶段通过A内扩展窗口选择A样本权重2，代理未来MSE由{phase_b['train_only_backtest_mse']:.6f}降至{phase_b['selected_backtest_mse']:.6f}。由于B标签不可见，该结果只表示代理回测改善。本文同时输出90%共形预测区间、工序重要性、漂移诊断、残差检验和高误差样本追溯结果。",
    )
    add_body(
        doc,
        "本文的主要贡献在于建立了折内可复现的数据治理流程，将工序结构和相对时间纳入特征表达，以扩展窗口评价替代单一随机切分，并明确区分A阶段、公开A答案后的B阶段以及事后评价的数据权限。该方法可为TFT-LCD制造过程的质量预警和工艺排查提供量化依据。",
    )
    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_before = Pt(8)
    run = keywords.add_run("关键词：")
    set_run_font(run, chinese="黑体", size=12, bold=True)
    run = keywords.add_run("TFT-LCD；机器学习；多工序制造；质量特性预测；数据漂移；XGBoost")
    set_run_font(run)

    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("ABSTRACT")
    set_run_font(run, chinese="黑体", size=16, bold=True)
    add_body(
        doc,
        f"TFT-LCD manufacturing involves hundreds of operations. Its production data are high-dimensional, small-sample, heterogeneous in scale, incomplete, and subject to batch drift. This study develops a machine-learning workflow for predicting the continuous product quality characteristic Y from equipment identifiers, process measurements, and production-time records. The Tianchi competition data contain {base['train_rows']} labeled training samples, {base['test_a_rows']} samples in test set A, and {base['test_b_rows']} samples in test set B, with 5,954 fields in the training file.",
        indent=False,
    )
    add_body(
        doc,
        "The workflow records column-level preprocessing decisions, imputes missing values using operation and tool groups with fold-local median fallback, and preserves missingness indicators. Timestamp fields are converted into operation duration and inter-operation gap features. Robust operation statistics are calculated after median and interquartile-range scaling. Feature selection combines ExtraTrees importance, target correlation, and temporal drift. Ridge regression, Random Forest, ExtraTrees, XGBoost, recency-weighted XGBoost, and a lightweight MLP ensemble are compared under random cross-validation, rolling validation, and temporal holdout evaluation.",
        indent=False,
    )
    add_body(
        doc,
        f"XGBoost achieved an MSE of 0.010063 under random out-of-fold evaluation but 0.032888 on the temporal holdout, demonstrating substantial optimism from random splitting. A training-only expanding-window ensemble reduced the post-hoc MSE on the published A labels from 0.030016 to {phase_a['selected_external_a_mse']:.6f}. After the A labels were released, an A-weighted three-member XGBoost ensemble reduced the pseudo-future MSE for stage B from {phase_b['train_only_backtest_mse']:.6f} to {phase_b['selected_backtest_mse']:.6f}. No B labels were used, so this result is not presented as the unknown leaderboard score. The system also reports conformal prediction intervals, operation importance, drift, residual diagnostics, and sample-level error traces.",
        indent=False,
    )
    keywords = doc.add_paragraph()
    run = keywords.add_run("KEY WORDS  ")
    set_run_font(run, chinese="黑体", size=12, bold=True)
    run = keywords.add_run("TFT-LCD; machine learning; multi-stage manufacturing; quality prediction; data drift; XGBoost")
    set_run_font(run)

    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("目  录")
    set_run_font(run, chinese="黑体", size=16, bold=True)
    toc_entries = [
        "1 绪论",
        "2 理论基础与相关研究",
        "3 数据审计与特征工程",
        "4 预测模型与实验设计",
        "5 实验结果与讨论",
        "6 结论与展望",
        "参考文献",
        "附录 复现与输出文件",
    ]
    for entry in toc_entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.2)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(entry)
        set_run_font(run, size=12)

    add_heading(doc, "1 绪论", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(
        doc,
        "制造过程的数字化使设备状态、工艺参数和质量检测结果能够以统一记录进入数据系统。智能制造研究据此将质量控制从生产后的抽样判定推进到生产过程中的预测与预警。Kusiak将智能制造概括为制造技术、信息技术和数据分析的协同应用，并指出数据驱动决策是其核心组成[3]。对于连续生产和复杂装配过程，质量预测模型能够在最终检测之前估计产品特性，从而支持异常隔离、工艺排查和检测资源配置。",
    )
    add_body(
        doc,
        "TFT-LCD生产包含多道成膜、曝光、刻蚀、清洗和检测工序。任何工序的设备状态或参数波动都可能影响最终电性和显示特性。天池智能制造质量预测赛题提供机台、温度、气体和液体流量、功率、制程时间等匿名变量，要求预测连续质量特性值Y[2]。该任务的样本数仅为数百量级，而字段接近六千，属于典型的小样本高维回归。字段还包含缺失、重复、常量、不同量纲、设备类别和时间信息，训练与未来批次之间也可能发生分布变化。",
    )
    add_body(
        doc,
        "生产现场需要模型同时满足准确、稳定和可追溯三个条件。只报告一次随机验证误差无法说明模型能否处理未来批次；只给出特征重要性也不能回答某个样本为何出现大误差。本文因此将数据治理、工序特征、批次验证、预测区间和样本诊断纳入同一流程。",
    )

    add_heading(doc, "1.2 研究意义", 2)
    add_body(
        doc,
        "在应用层面，连续质量特性的提前预测可以为工序放行提供辅助证据。当预测值偏离正常区间或预测区间明显变宽时，系统可以优先触发复检，而不是直接替代检测。对关键工序和设备组的误差分析还能帮助工程人员缩小排查范围。多工序质量预测研究已经表明，自适应模型和模型融合能够处理生产阶段之间的关联与工况变化[4-5]。",
    )
    add_body(
        doc,
        "在方法层面，本数据集揭示了随机交叉验证的局限。本文实验中，随机OOF误差约为0.010，而未来时间段误差接近0.033。若只选择随机CV最优模型，研究结论会高估部署性能。概念漂移研究指出，输入分布或输入与目标之间的关系会随时间变化[14]；预测研究也强调采用符合真实预测顺序的外推评价[15]。因此，本文重点研究适合小样本高维制造数据的时间泛化评价和增量适配。",
    )

    add_heading(doc, "1.3 研究内容", 2)
    add_body(
        doc,
        "本文首先建立字段、样本和工序三级数据审计，明确全空列、常量列、重复列、类别列与时间列的处理规则。随后利用工序与设备关系完成条件填充，以稳健标准化构造工序统计和相对时间特征，并联合模型贡献、相关性和批次漂移完成特征选择。模型实验比较线性模型、随机化树模型、梯度提升树和轻量神经网络，再通过滚动OOF优化非负融合权重。",
    )
    add_body(
        doc,
        "竞赛阶段采用两套独立策略。A阶段只使用原训练集，通过训练内扩展窗口选择三成员XGBoost；A答案仅在预测完成后用于事后评价。A答案公布后，B阶段把A标签作为新增监督数据，并在A内部再次构造扩展窗口选择样本权重。两阶段均保存基线、逐样本误差、策略汇总和数据权限清单。",
    )

    add_heading(doc, "1.4 技术路线与主要贡献", 2)
    add_body(
        doc,
        "本文技术路线依次为原始数据读取、结构审计、工序与设备解析、折内缺失处理、稳健工序统计、时间关系构造、稳定特征选择、多模型训练、三层验证、预测区间校准、A阶段模型选择、B阶段增量适配和结果解释。流程中的每一步都生成结构化文件或图形，便于复现和审核。",
    )
    add_body(
        doc,
        "本文的贡献包括四个方面。第一，缺失值不再统一填0，而是按工序设备组统计填充并保留缺失指示。第二，特征选择同时考虑预测贡献和时间漂移，工序统计在稳健标准化后计算。第三，验证设计同时保留随机CV、滚动CV和时间外推，A/B阶段采用符合竞赛顺序的扩展窗口。第四，模型输出包括共形区间、工序贡献、测试漂移和高误差样本证据，避免只交付一个无法解释的点预测。",
    )
    add_figure(doc, "15_pipeline_run_trace.png", "图1-1 基础预测流程运行轨迹")

    add_heading(doc, "2 理论基础与相关研究", 1)
    add_heading(doc, "2.1 智能制造质量预测", 2)
    add_body(
        doc,
        "质量预测利用生产过程变量估计尚未完成检测的产品质量。与一般表格回归相比，多工序制造数据存在阶段依赖、设备切换和批次漂移。Lughofer等研究了多阶段制造中的自适应质量监督与优化，强调模型应随过程状态变化调整[4]。向锋等针对复杂制造过程研究模型融合质量预测，说明不同模型的信息互补可提高复杂工况下的预测能力[5]。这些研究为本文的工序聚合、时间验证和模型融合提供了直接依据。",
    )
    add_body(
        doc,
        "屈方磊使用同一天池TFT-LCD数据开展Stacking质量预测研究，以GBR、XGBoost、RF和SVR为基学习器、线性回归为元学习器，并报告随机五折条件下Stacking MSE为0.01328[1]。该研究证明了同一数据上的集成可行性，但其流程采用全局Min-Max标准化、缺失填0和随机五折评价，未重点讨论工序结构、未来批次外推、预测区间和A/B标签权限。本文不重复该Stacking结构，而是把这些未充分处理的问题作为研究重点。",
    )
    add_table(
        doc,
        "表2-1 同数据集直接相关研究与本文的差异",
        ["比较项", "屈方磊研究", "本文"],
        [
            ["数据", "天池TFT-LCD复赛数据", "相同数据并区分A/B阶段"],
            ["缺失处理", "Min-Max后填0", "设备组均值 中位数兜底 缺失指示"],
            ["类别编码", "LabelEncoder", "未知类别安全One-Hot"],
            ["特征选择", "GBDT累计重要性90%", "贡献 相关性 漂移联合评分"],
            ["验证", "随机五折", "随机 滚动 时间外推 扩展窗口"],
            ["最终模型", "线性Stacking", "三种子XGBoost与阶段适配"],
            ["可信度", "MSE对比", "区间 漂移 残差 样本追溯"],
        ],
        widths=[3.0, 5.4, 7.0],
    )

    add_heading(doc, "2.2 树模型与集成学习", 2)
    add_body(
        doc,
        "梯度提升通过逐步拟合前一阶段的残差构造加法模型。Friedman给出了梯度提升机的函数优化解释[7]。XGBoost在此基础上引入二阶近似、正则化和高效并行实现，适合处理非线性、特征交互和高维表格数据[10]。随机森林通过样本和特征随机化降低树模型方差[8]，ExtraTrees进一步随机选择切分点以增加成员差异[9]。这些模型对单调变换和量纲差异相对不敏感，因此适合作为本文主干与对照。",
    )
    add_body(
        doc,
        "Stacking利用第一层模型的预测训练第二层学习器，是经典的泛化集成方法[12]。然而，小样本条件下若基模型OOF预测生成不严格，元学习器容易再次拟合噪声。本文采用非负且权重和为1的线性融合限制自由度，并通过滚动OOF确定权重。基础融合表达式如下。",
    )
    add_native_equation(doc, "ŷᵢ = Σₖ wₖ ŷᵢₖ，wₖ ≥ 0，Σₖ wₖ = 1", "2-1")
    add_body(
        doc,
        "Grinsztajn等在典型表格数据基准上发现，树模型在中等规模数据上仍常优于深度学习模型[13]。TabPFN进一步表明，表格基础模型在小数据条件下具有较强潜力[17]。考虑到当前训练样本只有800条，本文保留轻量MLP作为对照，而不预设深度模型必然优于树模型。",
    )

    add_heading(doc, "2.3 缺失数据和批次漂移", 2)
    add_body(
        doc,
        "缺失值处理取决于缺失机制。Little和Rubin系统区分了完全随机缺失、随机缺失和非随机缺失，并指出填充方法的有效性依赖缺失过程假设[6]。匿名制造数据无法仅凭数值表确认缺失原因，因此本文不把统计相关解释为物理因果，而是保留工序、设备和缺失指示证据，供后续与生产日志核对。",
    )
    add_body(
        doc,
        "批次漂移表现为输入分布、缺失模式或目标映射随生产阶段变化。Gama等总结了概念漂移检测与适配方法[14]。对本任务而言，时间戳并不一定形成完整、等间隔的序列，直接使用循环神经网络缺少可靠顺序依据。本文采用工序相对时间、滚动切分和近期样本权重，更贴合现有数据结构。",
    )

    add_heading(doc, "2.4 预测区间与残差检验", 2)
    add_body(
        doc,
        "点预测不能表达模型对单个产品的置信程度。共形预测利用校准残差构造有限样本预测区间，不要求预先指定误差正态分布[16]。本文使用滚动OOF或阶段扩展窗口残差计算90%区间半径。其区间形式为下式，其中q为校准绝对残差的经验分位数。",
    )
    add_native_equation(doc, "C₁₋α(x) = [ŷ(x) - q₁₋α，ŷ(x) + q₁₋α]", "2-2")
    add_body(
        doc,
        "残差检验包括Shapiro-Wilk正态性检验[18]、D'Agostino-Pearson正态性检验[19]、Breusch-Pagan异方差检验[20]和Ljung-Box自相关检验[21]。这些检验用于发现模型未解释的分布结构，不直接证明某个工艺因素造成误差。",
    )

    add_heading(doc, "3 数据审计与特征工程", 1)
    add_heading(doc, "3.1 数据来源与权限边界", 2)
    add_body(
        doc,
        f"数据来自天池智能制造质量预测复赛[2]。训练文件含{base['train_rows']}行和5954列，测试A含{base['test_a_rows']}行，测试B含{base['test_b_rows']}行。第一列为ID，最后一列为连续目标Y，其余字段按名称映射到210、220、300等多道工序。训练Y均值为{y_mean:.6f}，标准差为{y_std:.6f}，范围为{y_min:.6f}至{y_max:.6f}。",
    )
    add_body(
        doc,
        "A阶段只允许使用训练集标签。A答案公布后，A标签可以用于分析A误差和训练B模型，但不能反向进入A模型的参数选择。B标签不可见。代码分别保存A训练内策略清单和B阶段适配清单，使实验结论能够追溯到具体标签权限。",
    )
    add_table(
        doc,
        "表3-1 数据集与标签用途",
        ["数据集", "样本数", "标签状态", "用途"],
        [
            ["训练集", str(base["train_rows"]), "可用", "特征选择 模型训练 内部验证"],
            ["测试A", str(base["test_a_rows"]), "赛后公布", "A事后评价 B阶段新增监督"],
            ["测试B", str(base["test_b_rows"]), "不可见", "最终预测"],
        ],
        widths=[3.0, 2.5, 3.3, 6.6],
    )
    add_figure(doc, "02_target_and_temporal_split.png", "图3-1 目标分布与时间切分")

    add_heading(doc, "3.2 字段结构审计", 2)
    add_body(
        doc,
        f"程序首先识别全空、常量、重复、类别和时间字段。最终删除全空列{audit_counts.get('removed_all_nan', 0)}个、常量列{audit_counts.get('removed_constant', 0)}个、重复列{audit_counts.get('removed_duplicate', 0)}个；将{audit_counts.get('converted_timestamp', 0)}个字段转换为时间表示，保留{audit_counts.get('kept_numeric', 0)}个普通数值字段。`feature_audit.csv`对每列记录原位置、工序、缺失率、零值率、唯一值数量、时间戳比例和处理状态。",
    )
    add_body(
        doc,
        "删除规则以结构无信息为边界。全空列无法提供观测，常量列不含样本差异，完全重复列只增加计算和共线性。高相关但不完全重复的变量不按固定阈值批量删除，因为树模型可能利用局部差异，且不同工序的相似变量仍可能对应不同物理位置。",
    )
    add_figure(doc, "01_data_audit_by_operation.png", "图3-2 分工序字段审计结果")

    add_heading(doc, "3.3 缺失值处理", 2)
    add_body(
        doc,
        f"训练集中有{advanced['rows_with_any_missing']}条样本至少包含一个缺失值，单样本最高缺失率为{advanced['row_missing_rate_max']:.1%}。缺失率在字段和工序之间差异明显，因此统一填0会把“缺失”与真实零值混合。本文先依据字段所属工序匹配最近的设备类别列，在训练折内计算设备组均值；无法获得组均值时使用训练折中位数；对缺失率超过0.5%的入选字段另外生成缺失指示。",
    )
    add_body(
        doc,
        "填充器只在当前训练折拟合，验证折和测试集不参与统计量计算。这个顺序避免验证信息进入训练。缺失值表保留每一行的缺失数量、缺失率、全缺失工序、部分缺失工序和缺失最严重工序，使异常样本能够回查。",
    )
    add_figure(doc, "17_row_and_column_missingness.png", "图3-3 行列缺失率分布")
    add_figure(doc, "18_sample_operation_missing_heatmap.png", "图3-4 样本与工序缺失矩阵")

    add_heading(doc, "3.4 量纲处理和工序统计", 2)
    add_body(
        doc,
        "不同匿名字段可能分别表示温度、压力、流量、功率和时间，数值范围不能直接比较。本文对每个字段使用训练折中位数和四分位距构造稳健Z分数，并将绝对值截断到8以内。",
    )
    add_native_equation(doc, "zᵢⱼ = (xᵢⱼ - median(x·ⱼ)) / IQR(x·ⱼ)", "3-1")
    add_body(
        doc,
        "在标准化后的工序内计算均值、标准差、最小值、最大值、中位数、上下四分位数、范围和三倍IQR异常率；在原始数据上计算缺失率和零值率。这些统计压缩了同一工序的多测点信息，同时保留原始关键特征，避免只用工序均值丢失局部阈值。",
    )

    add_heading(doc, "3.5 时间与类别特征", 2)
    add_body(
        doc,
        "时间字段根据8位、14位和16位格式解析。每道工序生成开始日、结束日、中位日、持续分钟和有效率，相邻工序生成等待小时和是否重叠。时间原点只由训练折确定。该表示关注相对间隔和生产路线，不要求所有字段形成规则时间序列。",
    )
    add_body(
        doc,
        "设备字段采用One-Hot编码，验证集中的新设备类别自动映射为全零组合。与LabelEncoder相比，该方法不人为赋予设备编号大小关系。树模型主干无需全局归一化；Ridge和MLP分支在各自模型内部执行标准化。",
    )

    add_heading(doc, "3.6 稳定特征选择", 2)
    add_body(
        doc,
        "特征选择在每个训练折中独立完成。ExtraTrees给出非线性预测贡献，Pearson绝对相关反映边际线性关系，训练前后段中位数差除以IQR得到漂移量。三个量共同构成稳定分数。",
    )
    add_native_equation(doc, "sⱼ = (0.70 rⱼimp + 0.30 rⱼcorr) (1 / (1 + dⱼ))⁰·²⁵", "3-2")
    add_body(
        doc,
        "其中rimp和rcorr分别为贡献和相关性的百分位排名，d为截断后的漂移量。每折选择前320个原始数值字段。最终稳定性报告记录每个字段的入选率、平均贡献、相关性和漂移。方差只用于判断数据离散程度和异常量纲，不被直接解释为对Y的贡献。",
    )
    add_figure(doc, "19_variance_importance_drift.png", "图3-5 方差 预测贡献与漂移的关系")

    add_heading(doc, "3.7 理论分布与降维", 2)
    add_body(
        doc,
        f"本文对{tested_features}个稳定重要特征分别拟合正态、Laplace、Logistic和Student-t分布，先用AIC选出相对最优分布，再执行KS检验。AIC最优分布中有{consistent_features}个通过5%水平的一致性检验，因此不能把全部传感器统一假设为正态分布。",
    )
    add_body(
        doc,
        "降维接口支持不降维、PCA50、PCA95、SVD50、PLS20和Autoencoder16。统一Ridge滚动实验中，PCA50的MSE最低，为0.038703，但仍弱于树模型；Yeo-Johnson在混合量纲字段上出现数值失稳。因此降维保留为可插拔对照，不进入XGBoost主干。",
    )
    add_figure(doc, "20_theoretical_distribution_fit.png", "图3-6 重要特征与理论分布拟合")
    add_figure(doc, "21_transform_reducer_comparison.png", "图3-7 特征变换与降维对比")

    add_heading(doc, "4 预测模型与实验设计", 1)
    add_heading(doc, "4.1 评价指标", 2)
    add_body(
        doc,
        "赛题以均方误差作为主要指标。本文同时报告均方根误差、平均绝对误差和决定系数。MSE对大误差施加平方惩罚，适合识别少量严重偏差；MAE反映典型绝对误差；R²描述相对于均值预测的解释程度。",
    )
    add_native_equation(doc, "MSE = n⁻¹ Σᵢ(yᵢ - ŷᵢ)²，RMSE = √MSE", "4-1")

    add_heading(doc, "4.2 三层验证设计", 2)
    add_body(
        doc,
        f"第一层在训练开发段执行{base['validation_schemes'][0].replace('_', ' ')}五折，估计同分布条件下的拟合能力。第二层采用{base['rolling_calibration_rows']}个滚动OOF样本，只用较早样本预测随后批次，用于融合权重和共形区间校准。第三层锁定最后{base['temporal_holdout_rows']}条时间样本作为外推集，用于检查模型在未来阶段的误差。Tashman指出，预测精度评价应保持训练与未来观测的时间关系[15]。",
    )
    add_body(
        doc,
        "所有监督步骤都在训练折内拟合，包括设备组填充、编码、稳健尺度和特征选择。时间外推集不参与基础权重优化。该外推集曾用于识别验证设计问题，因此本文将其称为时间泛化基准，而不是完全未触碰的最终盲测。",
    )
    add_figure(doc, "04_oof_vs_temporal_holdout.png", "图4-1 随机OOF与时间外推误差")

    add_heading(doc, "4.3 候选模型与参数", 2)
    add_body(
        doc,
        "候选模型包括均值基线、Ridge、ExtraTrees、XGBoost、近期加权XGBoost和三成员MLP。RandomForest在高级分析中作为标准树模型对照。Ridge处理高维共线数据，其L2约束来源于Hoerl和Kennard提出的岭回归[11]。XGBoost主干使用1200棵树、学习率0.05、最大深度2、最小子节点权重3、样本比例0.80和特征比例0.55。嵌套网格搜索表明最小子节点权重3在多数外层折更合适。",
    )
    add_body(
        doc,
        "近期加权模型按样本阶段排名施加指数衰减，使接近验证时点的数据具有更高权重。基础融合权重仅由滚动OOF通过约束优化获得。最终权重约为XGBoost 60.6%、近期加权XGBoost 33.9%和ExtraTrees 5.5%；Ridge和MLP权重为0，因此只保留为对照。",
    )
    add_figure(doc, "05_blend_weights.png", "图4-2 基础模型非负融合权重")

    add_heading(doc, "4.4 消融与嵌套搜索", 2)
    add_body(
        doc,
        "特征消融依次比较原始特征、原始加缺失指示、工序统计、工序时间、完整类别视图和纯工序聚合。三折实验中工序时间视图MSE为0.027388，但在四折主流程中优势未稳定复现，因此最终保留完整视图。嵌套搜索将参数选择限制在外层训练折，外层时间MSE为0.027636，防止使用同一验证结果反复调参。",
    )
    add_figure(doc, "23_feature_ablation.png", "图4-3 工序特征消融实验")
    add_figure(doc, "24_nested_grid_search.png", "图4-4 嵌套网格搜索结果")

    add_heading(doc, "4.5 A阶段与B阶段实验协议", 2)
    add_body(
        doc,
        f"A阶段按竞赛样本顺序设置三个扩展窗口：前500、600和700条分别预测随后100条，共{phase_a['validation_rows']}个伪未来样本。比较均匀权重和半衰期0.20、0.35、0.50的近期加权XGBoost，每个策略对{phase_a['ensemble_members']}个独立随机种子的特征选择和模型预测取平均。选定策略并完成A预测后，程序才读取公开A答案。",
    )
    add_body(
        doc,
        f"B阶段允许使用已经公布的A标签。以A前100、150和200条作为新增监督数据，分别预测其后50、50和100条，共{phase_b['validation_rows']}个伪未来样本。候选A样本权重为1、2和4，每种策略使用{phase_b['ensemble_members']}成员平均。程序根据合并MSE锁定权重，再使用800条原训练数据和300条A标签拟合B模型。B标签从未参与选择或训练。",
    )
    add_table(
        doc,
        "表4-1 A阶段与B阶段验证协议",
        ["阶段", "训练标签", "扩展窗口", "最终训练", "禁止使用"],
        [
            ["A", "原训练集", "前500 600 700条各预测100条", "800条训练标签", "A标签"],
            ["B", "原训练集与A前缀", "A前100预测50 前150预测50 前200预测100", "800条训练与300条A", "B标签"],
        ],
        widths=[2.0, 3.6, 4.8, 3.6, 2.0],
    )

    add_heading(doc, "5 实验结果与讨论", 1)
    add_heading(doc, "5.1 基础模型结果", 2)
    add_table(
        doc,
        "表5-1 基础模型在三种验证口径下的结果",
        ["模型", "随机OOF MSE", "滚动OOF MSE", "时间外推MSE", "外推R²"],
        chapter_model_table(model_summary),
        widths=[3.5, 3.0, 3.0, 3.0, 2.5],
    )
    add_body(
        doc,
        "XGBoost在随机OOF和时间外推中均为最强单模型。RobustBlend的滚动OOF MSE为0.025261，略低于普通XGBoost的0.025421；时间外推MSE为0.032885，与XGBoost的0.032888接近。ExtraTrees能够提供少量多样性，但独立外推误差较高。Ridge和MLP在跨批次预测中不稳定，因而没有进入最终基础融合。",
    )
    add_body(
        doc,
        "随机OOF与外推结果之间的差距是本研究最重要的验证发现。随机切分把相邻批次分散到训练折和验证折，模型可以利用相似工况；未来批次则可能同时改变设备组合、缺失模式和目标均值。后续A/B阶段因此使用扩展窗口，而不根据随机CV单独决定提交模型。",
    )
    add_figure(doc, "03_cv_model_stability.png", "图5-1 候选模型交叉验证稳定性")
    add_figure(doc, "07_temporal_holdout_prediction.png", "图5-2 时间外推真实值与预测值")

    add_heading(doc, "5.2 树模型和降维对比", 2)
    tree = pd.read_csv(OUTPUT / "tree_model_comparison_summary.csv")
    add_table(
        doc,
        "表5-2 标准树模型滚动对比",
        ["模型", "MSE", "RMSE", "MAE", "R²"],
        [
            [
                str(row.model),
                f"{row.mse_mean:.6f}",
                f"{row.rmse_mean:.6f}",
                f"{row.mae_mean:.6f}",
                f"{row.r2_mean:.4f}",
            ]
            for row in tree.itertuples(index=False)
        ],
        widths=[4.0, 3.0, 3.0, 3.0, 2.5],
    )
    add_body(
        doc,
        "标准树模型对比中XGBoost误差最低，说明浅层提升树能够在当前样本量下利用关键传感器的非线性交互。RandomForest和ExtraTrees更强调方差降低，但其平均化也会压缩极端预测。PCA50虽然改善了线性模型，却没有达到树模型误差，因此没有必要先把所有原始工序压缩到统一线性空间。",
    )
    add_figure(doc, "25_tree_model_comparison.png", "图5-3 RandomForest ExtraTrees与XGBoost对比")

    phase_a_heading = add_heading(doc, "5.3 A阶段结果", 2)
    phase_a_heading.paragraph_format.page_break_before = True
    phase_a_summary = pd.read_csv(OUTPUT / "phase_a_strategy_summary.csv")
    add_table(
        doc,
        "表5-3 A阶段训练内策略选择",
        ["策略", "回测MSE", "RMSE", "MAE", "最差折MSE"],
        [
            [
                str(row.strategy),
                f"{row.mse:.6f}",
                f"{row.rmse:.6f}",
                f"{row.mae:.6f}",
                f"{row.worst_fold_mse:.6f}",
            ]
            for row in phase_a_summary.itertuples(index=False)
        ],
        widths=[5.0, 2.8, 2.8, 2.8, 3.0],
    )
    add_body(
        doc,
        f"训练内回测选择均匀权重XGBoost，MSE为{phase_a['selected_backtest_mse']:.6f}。近期加权在第一窗口有时有效，但在后续窗口没有稳定优势。策略锁定后的A真实评价显示，三成员均匀权重模型MSE为{phase_a['selected_external_a_mse']:.6f}、RMSE为{phase_a['selected_external_a_rmse']:.6f}、MAE为{phase_a['selected_external_a_mae']:.6f}、R²为{phase_a['selected_external_a_r2']:.4f}。相对于RobustBlend基线，MSE下降{phase_a['external_a_relative_mse_improvement']:.1%}。",
    )
    add_body(
        doc,
        "公开A答案显示模型仍存在约0.050的正偏差，预测标准差也小于真实标准差，说明极端值收缩尚未完全消除。该结果用于评估已经锁定的A模型，不用于重新挑选半衰期。按公开答案选择表现更好的候选会使A结果带有标签泄漏。",
    )
    add_figure(doc, "30_phase_a_training_only_backtest.png", "图5-4 A阶段训练内选择与事后评价")

    phase_b_heading = add_heading(doc, "5.4 B阶段适配结果", 2)
    phase_b_heading.paragraph_format.page_break_before = True
    phase_b_summary = pd.read_csv(OUTPUT / "phase_b_strategy_summary.csv")
    add_table(
        doc,
        "表5-4 B阶段公开A标签增量训练回测",
        ["策略", "A权重", "回测MSE", "MAE", "最差折MSE"],
        [
            [
                str(row.strategy),
                f"{row.a_weight:g}",
                f"{row.mse:.6f}",
                f"{row.mae:.6f}",
                f"{row.worst_fold_mse:.6f}",
            ]
            for row in phase_b_summary.itertuples(index=False)
        ],
        widths=[6.0, 2.2, 2.8, 2.8, 3.2],
    )
    add_body(
        doc,
        f"三成员回测选择A样本权重2。适配模型MSE为{phase_b['selected_backtest_mse']:.6f}，相对于A阶段训练集-only模型在相同200个样本上的{phase_b['train_only_backtest_mse']:.6f}下降{phase_b['relative_mse_improvement']:.1%}。适配后RMSE为{phase_b['selected_backtest_rmse']:.6f}、MAE为{phase_b['selected_backtest_mae']:.6f}、R²为{phase_b['selected_backtest_r2']:.4f}，平均偏差接近零。",
    )
    add_body(
        doc,
        f"最终B预测均值为{phase_b['b_prediction_mean']:.6f}，标准差为{phase_b['b_prediction_std']:.6f}。与不使用A标签的B基线相比，平均绝对调整量为{phase_b['mean_absolute_change_vs_train_only']:.6f}。这些变化与A阶段发现的整体高估和方差收缩方向一致，但B真实标签不可见，因此不能把29.1%的代理回测下降写成B榜真实提升。",
    )
    add_figure(doc, "29_phase_b_adaptation_backtest.png", "图5-5 B阶段公开A标签适配回测")

    add_heading(doc, "5.5 工序贡献与分布漂移", 2)
    add_body(
        doc,
        f"跨折贡献最高的工序为{'、'.join(top_operations)}。训练集与A/B测试集漂移较高的工序为{'、'.join(drift_operations)}。工序344、520等同时出现在贡献或漂移前列，应优先检查设备组合、缺失率和参数范围。但重要性表示模型使用频率，漂移表示统计差异，两者都不能单独证明工序对Y具有因果作用。",
    )
    add_figure(doc, "11_operation_importance.png", "图5-6 关键工序预测贡献")
    add_figure(doc, "12_operation_distribution_drift.png", "图5-7 工序分布漂移")

    add_heading(doc, "5.6 残差和高误差样本", 2)
    residual_rows = []
    for row in residual_tests.itertuples(index=False):
        conclusion = "显著" if float(row.pvalue) < 0.05 else "不显著"
        residual_rows.append([str(row.test), f"{float(row.statistic):.4f}", f"{float(row.pvalue):.4g}", conclusion])
    add_table(
        doc,
        "表5-5 时间外推残差统计检验",
        ["检验", "统计量", "p值", "5%结论"],
        residual_rows,
        widths=[7.0, 3.0, 3.0, 2.5],
    )
    add_body(
        doc,
        "Shapiro-Wilk和D'Agostino-Pearson检验拒绝残差正态性，Ljung-Box在滞后5和10阶显著，说明残差仍含生产顺序结构。Breusch-Pagan检验未发现显著的预测值相关异方差。非正态和自相关支持采用经验共形区间，并提示未来可加入批次状态或滞后工序信息。",
    )
    add_body(
        doc,
        f"时间外推最大误差样本为{advanced['highest_error_sample']}。该样本单点对平方误差影响较大，项目已输出其关键特征稳健Z分数和所属工序。没有生产日志证明其标签或采集错误，因此评价时保留该样本，只进行敏感性分析。",
    )
    add_figure(doc, "27_residual_qq_and_autocorrelation.png", "图5-8 残差Q-Q图与生产顺序自相关")
    add_figure(doc, "26_high_error_sample_diagnosis.png", "图5-9 高误差样本局部诊断")

    add_heading(doc, "5.7 预测区间与部署讨论", 2)
    add_body(
        doc,
        f"基础模型的90%共形区间在时间外推集上实际覆盖率为{base['holdout_coverage']:.1%}，区间半径为{base['conformal_radius']:.4f}。A阶段和B阶段使用各自扩展窗口残差重新校准区间，避免沿用不同数据权限下的误差分布。覆盖率只在已有验证样本上成立；当设备或材料发生明显变化时，需要重新检查漂移和覆盖率。",
    )
    add_figure(doc, "13_conformal_interval_coverage.png", "图5-10 共形预测区间覆盖率")
    add_body(
        doc,
        "部署时可将模型输出分为点预测、预测区间、漂移分数和关键工序四部分。区间完全落在工艺正常范围内的样本可降低抽检优先级，区间越界或漂移较大的样本应优先复检。模型不应直接自动修改设备参数，因为匿名数据无法确认字段单位、可控性和安全边界。工艺调整仍需工程人员审批和小范围验证。",
    )

    add_heading(doc, "6 结论与展望", 1)
    add_heading(doc, "6.1 研究结论", 2)
    add_body(
        doc,
        "本文针对TFT-LCD小样本高维多工序数据建立了完整的质量特性预测流程。该流程从字段审计开始，在训练折内完成工序设备条件填充、稳健统计、时间关系构造和抗漂移选择，并对线性模型、随机化树模型、梯度提升树和轻量神经网络进行统一比较。实验表明XGBoost仍是最可靠的主干，复杂Stacking或无约束线性分支并未在未来批次中稳定获益。",
    )
    add_body(
        doc,
        "验证结果说明，随机五折不能代表批次外推性能。随机OOF MSE约为0.010，而时间外推MSE约为0.033。采用扩展窗口和三随机种子平均后，A模型在不使用A标签选择或拟合的前提下取得0.027525的真实MSE。A答案公布后，权重为2的增量训练在B代理未来样本上取得0.020916的MSE。项目同时保留基线版本和逐样本证据，使每项改动都能追溯。",
    )
    add_body(
        doc,
        "本文没有把特征重要性解释为因果效应，也没有把B代理回测写成未知B榜得分。共形区间、漂移报告和高误差样本分析为模型进入质量预警流程提供了必要的风险信息。",
    )

    add_heading(doc, "6.2 研究局限", 2)
    add_body(
        doc,
        "第一，字段经过匿名化，缺少单位、设备维护、材料批次和工艺上下限，统计关联无法转换为明确的物理机理。第二，训练样本只有800条，极端质量样本较少，模型对分布尾部仍存在收缩。第三，时间字段并不构成完整的等间隔序列，本文的相对时间特征不能替代真实工艺路线图。第四，B标签不可见，B阶段只完成代理验证，真实效果需要平台或新增生产周期确认。",
    )

    add_heading(doc, "6.3 后续研究", 2)
    add_body(
        doc,
        "后续研究应优先获取字段字典、设备事件、材料批次和维护记录，以验证缺失和漂移的真实原因。模型方面可研究按工序分块的监督降维、面向极端样本的分位数损失、在线漂移检测和带遗忘机制的增量学习。若引入TabPFN等表格基础模型，应在相同扩展窗口、相同特征权限和相同计算预算下比较，不能只报告随机划分结果。",
    )
    add_body(
        doc,
        "在实际部署前还需要建立回退策略。模型输入超出训练范围、区间覆盖率下降或设备类别大量新增时，应自动提高抽检比例并暂停自动放行。只有在连续生产周期中验证误报率、漏报率和维护成本后，才能讨论减少检测工序。",
    )

    add_heading(doc, "参考文献", 1)
    references = [
        "[1] 屈方磊. 基于集成学习的智能制造质量预测研究[D]. 中北大学, 2024.",
        "[2] 阿里云天池. 智能制造质量预测赛题[EB/OL]. https://tianchi.aliyun.com/competition/entrance/231633/introduction, 访问日期: 2026-09-13.",
        "[3] KUSIAK A. Smart manufacturing[J]. International Journal of Production Research, 2018, 56(1-2): 508-517. DOI: 10.1080/00207543.2017.1351644.",
        "[4] LUGHOFER E, ZAVOIANU A C, POLLAK R, et al. Autonomous supervision and optimization of product quality in a multi-stage manufacturing process based on self-adaptive prediction models[J]. Journal of Process Control, 2019, 76: 27-45. DOI: 10.1016/j.jprocont.2019.02.005.",
        "[5] XIANG F, YANG L, ZHANG M, et al. Model fusion based product quality prediction for complex manufacturing process[J]. Scientia Sinica Technologica, 2023, 53: 1127-1137. DOI: 10.1360/SST-2022-0427.",
        "[6] LITTLE R J A, RUBIN D B. Statistical Analysis with Missing Data[M]. 2nd ed. Hoboken: Wiley, 2002. DOI: 10.1002/9781119013563.",
        "[7] FRIEDMAN J H. Greedy function approximation: A gradient boosting machine[J]. The Annals of Statistics, 2001, 29(5): 1189-1232. DOI: 10.1214/aos/1013203451.",
        "[8] BREIMAN L. Random forests[J]. Machine Learning, 2001, 45(1): 5-32. DOI: 10.1023/A:1010933404324.",
        "[9] GEURTS P, ERNST D, WEHENKEL L. Extremely randomized trees[J]. Machine Learning, 2006, 63(1): 3-42. DOI: 10.1007/s10994-006-6226-1.",
        "[10] CHEN T, GUESTRIN C. XGBoost: A scalable tree boosting system[C]//Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining. New York: ACM, 2016: 785-794. DOI: 10.1145/2939672.2939785.",
        "[11] HOERL A E, KENNARD R W. Ridge regression: Biased estimation for nonorthogonal problems[J]. Technometrics, 1970, 12(1): 55-67. DOI: 10.1080/00401706.1970.10488634.",
        "[12] WOLPERT D H. Stacked generalization[J]. Neural Networks, 1992, 5(2): 241-259. DOI: 10.1016/S0893-6080(05)80023-1.",
        "[13] GRINSZTAJN L, OYALLON E, VAROQUAUX G. Why do tree-based models still outperform deep learning on typical tabular data?[C]//Advances in Neural Information Processing Systems 35. 2022: 507-520. DOI: 10.52202/068431-0037.",
        "[14] GAMA J, ZLIOBAITE I, BIFET A, et al. A survey on concept drift adaptation[J]. ACM Computing Surveys, 2014, 46(4): 1-37. DOI: 10.1145/2523813.",
        "[15] TASHMAN L J. Out-of-sample tests of forecasting accuracy: An analysis and review[J]. International Journal of Forecasting, 2000, 16(4): 437-450. DOI: 10.1016/S0169-2070(00)00065-0.",
        "[16] ANGELOPOULOS A N, BATES S. Conformal prediction: A gentle introduction[J]. Foundations and Trends in Machine Learning, 2023, 16(4): 494-591. DOI: 10.1561/2200000101.",
        "[17] HOLLMANN N, MULLER S, PURUCKER L, et al. Accurate predictions on small data with a tabular foundation model[J]. Nature, 2025, 637(8045): 319-326. DOI: 10.1038/s41586-024-08328-6.",
        "[18] SHAPIRO S S, WILK M B. An analysis of variance test for normality complete samples[J]. Biometrika, 1965, 52(3-4): 591-611. DOI: 10.2307/2333709.",
        "[19] D'AGOSTINO R, PEARSON E S. Tests for departure from normality empirical results for the distributions of b2 and square root of b1[J]. Biometrika, 1973, 60(3): 613-622. DOI: 10.1093/biomet/60.3.613.",
        "[20] BREUSCH T S, PAGAN A R. A simple test for heteroscedasticity and random coefficient variation[J]. Econometrica, 1979, 47(5): 1287-1294. DOI: 10.2307/1911963.",
        "[21] LJUNG G M, BOX G E P. On a measure of lack of fit in time series models[J]. Biometrika, 1978, 65(2): 297-303. DOI: 10.1093/biomet/65.2.297.",
    ]
    for reference_index, reference in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        if reference_index == 13:
            paragraph.paragraph_format.page_break_before = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)
        paragraph.paragraph_format.line_spacing = 1.18
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(reference)
        set_run_font(run, size=10)

    add_heading(doc, "附录 复现与输出文件", 1)
    add_body(
        doc,
        "项目完整入口为python -m innovative_solution.run。主流程生成字段审计、折级指标、融合权重、逐样本预测、工序贡献、漂移、共形区间和提交文件；高级分析生成缺失、分布、降维、消融、网格搜索和残差证据；A/B阶段分别生成独立回测清单。所有CSV均保存明确列名，竞赛提交文件除外。",
    )
    output_rows = [
        ["submission_A.csv", "A阶段正式两列无表头提交"],
        ["submission_B.csv", "B阶段正式两列无表头提交"],
        ["feature_audit.csv", "逐字段处理理由"],
        ["row_missing_report.csv", "逐样本缺失追溯"],
        ["model_summary.csv", "三种验证口径模型指标"],
        ["phase_a_manifest.json", "A阶段选择和标签权限"],
        ["phase_b_manifest.json", "B阶段适配和无B标签声明"],
        ["WORKFLOW_DESCRIPTION.md", "完整工作过程文字稿"],
    ]
    add_table(doc, "表A-1 核心复现文件", ["文件", "内容"], output_rows, widths=[6.0, 9.4])
    add_body(
        doc,
        "当前两份正式提交均已核验：A文件300行，B文件412行，均为两列、无表头、ID顺序与测试文件完全一致，预测值不存在缺失或无穷。基础测试覆盖时间解析、近期权重、非负融合、共形分位数、A答案ID顺序和降维接口。",
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
